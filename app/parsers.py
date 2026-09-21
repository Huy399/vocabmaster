"""
File parsers for PDF and DOCX documents with text extraction and context matching.
"""
import io
import re
import unicodedata
import logging
from typing import List, Optional
import pypdf
import docx

logger = logging.getLogger(__name__)

try:
    import pymupdf
    HAS_PYMUPDF = True
except ImportError:
    HAS_PYMUPDF = False

# Biểu thức chính quy phát hiện tiếng Việt bao gồm cả Unicode NFC dựng sẵn và NFD tổ hợp
VN_ACCENT_PATTERN = re.compile(
    r'[àáạảãâầấậẩẫăằắặẳẵèéẹẻẽêềếệểễìíịỉĩòóọỏõôồốộổỗơờớợởỡùúụủũưừứựửữỳýỵỷỹđ'
    r'\u0300\u0301\u0303\u0309\u0323\u02c6\u0306\u031b]',
    re.IGNORECASE
)

def normalize_unicode(text: str) -> str:
    """
    Chuẩn hóa toàn bộ chuỗi về Unicode NFC chuẩn quốc tế và loại bỏ khoảng trắng vô hình:
    - Loại bỏ Zero-width space (\u200b), BOM (\ufeff), Soft hyphen (\u00ad).
    - Chuẩn hóa Non-breaking space (\u00a0, \u202f, \u3000) về khoảng trắng đơn chuẩn.
    """
    if not text:
        return ""
    normalized = unicodedata.normalize("NFC", text)
    normalized = re.sub(r'[\u200b\u200c\u200d\u200e\u200f\ufeff\u00ad]', '', normalized)
    normalized = normalized.replace('\u00a0', ' ').replace('\u202f', ' ').replace('\u3000', ' ')
    return normalized

def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Trích xuất văn bản từ nội dung nhị phân của file PDF bằng PyMuPDF hoặc pypdf fallback an toàn."""
    if not file_bytes or len(file_bytes) < 8:
        return ""

    extracted_chunks = []
    if HAS_PYMUPDF:
        try:
            with pymupdf.open(stream=file_bytes, filetype="pdf") as doc:
                if doc.is_encrypted:
                    try:
                        doc.authenticate("")
                    except Exception as auth_err:
                        logger.debug(f"PyMuPDF authenticate failed: {auth_err}")
                for page in doc:
                    try:
                        tabs = page.find_tables()
                        if tabs and tabs.tables:
                            for tab in tabs.tables:
                                for row in tab.extract():
                                    clean_cells = [
                                        str(c).strip().replace('\n', ' ')
                                        for c in row
                                        if c is not None and str(c).strip() and str(c).strip().lower() != 'none'
                                    ]
                                    if clean_cells:
                                        extracted_chunks.append(" | ".join(clean_cells))
                    except Exception as tab_err:
                        logger.debug(f"PyMuPDF page table extraction failed: {tab_err}")
                    try:
                        t = page.get_text("text")
                        if t and t.strip():
                            extracted_chunks.append(t.strip())
                    except Exception as text_err:
                        logger.debug(f"PyMuPDF page text extraction failed: {text_err}")
            if extracted_chunks:
                return normalize_unicode("\n\n".join(extracted_chunks))
        except Exception as e:
            logger.warning(f"PyMuPDF read error: {e}, falling back to pypdf")

    # Fallback an toàn sang pypdf
    try:
        reader = pypdf.PdfReader(io.BytesIO(file_bytes))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as dec_err:
                logger.debug(f"pypdf decrypt failed: {dec_err}")
        for page_idx, page in enumerate(reader.pages):
            try:
                page_text = page.extract_text()
                if page_text and page_text.strip():
                    extracted_chunks.append(page_text.strip())
            except Exception as e:
                logger.error(f"Error reading PDF page {page_idx}: {e}")
    except Exception as e:
        logger.error(f"pypdf error: {e}", exc_info=True)

    return normalize_unicode("\n\n".join(extracted_chunks))

def extract_text_from_docx(file_bytes: bytes) -> str:
    """Trích xuất văn bản từ DOCX bảo toàn thứ tự xuất hiện giữa đoạn văn và bảng biểu."""
    if not file_bytes or len(file_bytes) < 8:
        return ""
    try:
        doc = docx.Document(io.BytesIO(file_bytes))
    except Exception as e:
        logger.error(f"Failed to open DOCX: {e}", exc_info=True)
        return ""

    chunks = []
    try:
        for element in doc.element.body:
            if element.tag.endswith('p'):
                text = "".join(node.text for node in element.iter() if node.tag.endswith('t') and node.text).strip()
                if text:
                    chunks.append(text)
            elif element.tag.endswith('tbl'):
                table = docx.table.Table(element, doc)
                for row in table.rows:
                    seen_cells = set()
                    row_cells = []
                    for cell in row.cells:
                        if cell._tc not in seen_cells:
                            seen_cells.add(cell._tc)
                            c_text = cell.text.strip().replace('\n', ' ')
                            if c_text:
                                row_cells.append(c_text)
                    if row_cells:
                        chunks.append(" | ".join(row_cells))
            elif element.tag.endswith('sdt'):
                for p in element.iter():
                    if p.tag.endswith('p'):
                        p_text = "".join(node.text for node in p.iter() if node.tag.endswith('t') and node.text).strip()
                        if p_text:
                            chunks.append(p_text)
    except Exception as e:
        logger.warning(f"DOCX body element traversal failed: {e}; falling back to paragraphs/tables")
        for p in doc.paragraphs:
            if p.text.strip():
                chunks.append(p.text.strip())
        for table in doc.tables:
            for row in table.rows:
                r_text = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                if r_text:
                    chunks.append(r_text)

    return normalize_unicode("\n\n".join(chunks))

def extract_text_from_txt(file_bytes: bytes) -> str:
    """Tự động nhận diện bảng mã UTF-8-SIG, UTF-8, UTF-16, Windows-1258, Latin-1."""
    if not file_bytes:
        return ""
    encodings = ["utf-8-sig", "utf-8", "utf-16", "windows-1258", "cp1252", "latin-1"]
    for enc in encodings:
        try:
            return normalize_unicode(file_bytes.decode(enc).strip())
        except UnicodeDecodeError:
            continue
    return normalize_unicode(file_bytes.decode("utf-8", errors="ignore").strip())

def clean_text(raw_text: str) -> str:
    """Làm sạch văn bản, loại bỏ header/footer rác, website và chuẩn hóa khoảng trắng."""
    normalized = normalize_unicode(raw_text)
    # Chỉ loại bỏ các liên kết quảng cáo thường gặp trong tài liệu luyện thi (giữ lại URL nội dung)
    cleaned = re.sub(r'https?://(?:www\.)?(?:tailieuonthi|facebook\.com|zalo\.me|citienglish)\S*', '', normalized, flags=re.IGNORECASE)
    cleaned = re.sub(r'TAILIEUONTHI\.NET', '', cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r'Tài Liệu Ôn Thi Group', '', cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r'TRUNG TÂM NGOẠI NGỮ CITI ENGLISH', '', cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r'HOTLINE:\s*[\d\s]+', '', cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r'Facebook:\s*\S+', '', cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r'Compiled by\s+[^\n]+', '', cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r'Tiếng Anh Mr\.\s*Cường\s*\d*', '', cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r'TỪ VỰNG TIẾNG ANH SONG NGỮ BỔ SUNG', '', cleaned, flags=re.IGNORECASE)
    # Loại bỏ các đoạn script / style / iframe / svg độc hại khỏi văn bản bóc tách
    cleaned = re.sub(r'(?is)<\s*(?:script|style|svg|iframe|object|embed)[^>]*>.*?<\s*/\s*(?:script|style|svg|iframe|object|embed)\s*>', '', cleaned)
    cleaned = re.sub(r'(?is)<\s*(?:script|style|svg|iframe|object|embed)[^>]*>', '', cleaned)
    
    cleaned = re.sub(r'[\r\t]+', ' ', cleaned)
    cleaned = re.sub(r'\n{3,}', '\n\n', cleaned)
    cleaned = re.sub(r'[ ]{2,}', ' ', cleaned)
    return cleaned.strip()

def split_into_sentences(text: str) -> List[str]:
    """
    Tách văn bản thành danh sách các câu tiếng Anh hoàn chỉnh, mạch lạc và sạch sẽ:
    - Loại bỏ phiên âm IPA, nhãn Synonym, Antonym, nhãn từ loại
    - Bóc tách câu ví dụ chuẩn từ các khối mục từ vựng (bilingual vocabulary book)
    - Loại bỏ tiếng Việt chú thích và các ký tự rác
    """
    raw_clean = clean_text(text)
    clean_sentences = set()

    # 0. Bảo vệ tuyệt đối URLs khỏi các biểu thức regex phiên âm và ký hiệu
    url_map = {}
    def mask_url(match):
        token = f"URLPROTECTEDTOKEN{len(url_map)}"
        url_map[token] = match.group(0)
        return token
    raw_clean = re.sub(r'https?://[^\s<>"]+|www\.[^\s<>"]+', mask_url, raw_clean)
    
    # 1. Bóc tách câu từ các khối từ vựng có đánh số (ví dụ: 26 BIOME (n) ...)
    blocks = re.split(r'\n(?=\s*\d+\s+[A-Z\s-]{2,}\s*\([a-z,.\s]+\))', raw_clean)
    for b in blocks:
        b_sub = re.sub(r'^\s*\d+\s+[A-Z\s-]{2,}\s*\([a-z,.\s]+\)', '', b)
        b_sub = re.sub(r'/[^/\n]+/', '', b_sub)
        b_sub = re.sub(r'^\s*(Synonym|Antonym|Synonyms|Antonyms)\s*:[^\n]*', '', b_sub, flags=re.MULTILINE | re.IGNORECASE)
        
        lines = [l.strip() for l in b_sub.split('\n') if l.strip()]
        eng_lines = []
        for l in lines:
            vn_chars = len(VN_ACCENT_PATTERN.findall(l))
            if vn_chars == 0 and not l.startswith(('Synonym', 'Antonym')) and not re.match(r'^\d+$', l):
                eng_lines.append(l)
        joined = ' '.join(eng_lines)
        sents = re.findall(r'([A-Z][^\n]{20,}[.!?])', joined)
        for s in sents:
            s_strip = s.strip()
            if ':' not in s_strip and not re.search(r'\((n|v|adj|adv)\)', s_strip):
                clean_sentences.add(s_strip)

    # 2. Bóc tách câu từ văn bản thông thường
    text_for_body = re.sub(r'/[a-zəɪɒʌʊæeɛθðʃʒŋˈˌːˑ.ˈˌ\-ˌɔɑɜɚɹɾɫɬɮβçɧħʔʕɦɥʋɰɻʁʀɢɴɱɲɳʈɖɟɡᵻᵿ]+/', '', raw_clean, flags=re.IGNORECASE)
    text_for_body = re.sub(r'^\s*(Synonym|Antonym|Synonyms|Antonyms)\s*:[^\n]*', '', text_for_body, flags=re.MULTILINE | re.IGNORECASE)
    text_for_body = re.sub(r'\([^\)]*[àáạảãâầấậẩẫăằắặẳẵèéẹẻẽêềếệểễìíịỉĩòóọỏõôồốộổỗơờớợởỡùúụủũưừứựửữỳýỵỷỹđ][^\)]*\)', '', text_for_body)
    
    paragraphs = text_for_body.split('\n\n')
    for p in paragraphs:
        p_clean = re.sub(r'\s*\n\s*', ' ', p).strip()
        raw_sents = re.split(r'(?<=[.?!])\s+', p_clean)
        for s in raw_sents:
            s_clean = re.sub(r'\s+', ' ', s).strip()
            s_clean = re.sub(r'^[)\],.:;\s]+', '', s_clean).strip()
            if re.match(r'^[A-Z]', s_clean) and s_clean.endswith(('.', '!', '?')) and 20 <= len(s_clean) <= 250:
                vn_count = len(VN_ACCENT_PATTERN.findall(s_clean))
                if vn_count <= 1 and ':' not in s_clean and not re.search(r'\b(Synonym|Antonym)\b', s_clean, re.IGNORECASE) and not re.search(r'\((n|v|adj|adv)\)', s_clean):
                    clean_sentences.add(s_clean)

    def sentence_quality_score(s: str) -> int:
        length = len(s)
        if 50 <= length <= 130:
            return 1000 - abs(length - 80)
        elif 30 <= length < 50:
            return 500 + length
        else:
            return 200 - length

    # Khôi phục nguyên vẹn các URLs
    restored_sentences = []
    for s in clean_sentences:
        for token, orig_url in url_map.items():
            s = s.replace(token, orig_url)
        restored_sentences.append(s)

    return sorted(restored_sentences, key=sentence_quality_score, reverse=True)

def find_context_sentence(word: str, sentences: List[str]) -> Optional[str]:
    """Tìm câu trong tài liệu có chứa từ vựng tương ứng để làm ngữ cảnh thực tế."""
    w = (word or "").strip()
    if not w:
        return None
    pattern = re.compile(rf'\b{re.escape(w)}\b', re.IGNORECASE)
    for s in sentences:
        if pattern.search(s):
            return s
    return None

def build_word_context_map(sentences: List[str]) -> dict:
    """Tạo bản đồ từ -> câu ngữ cảnh O(1). Ưu tiên các câu ngắn gọn, mạch lạc."""
    word_to_sentence = {}
    for s in sentences:
        words = re.findall(r'\b[a-zA-Z]{3,}\b', s.lower())
        # Hỗ trợ cả các từ ghép nối bằng dấu gạch ngang (compound words)
        compounds = re.findall(r'\b[a-zA-Z]{2,}(?:-[a-zA-Z]{2,})+\b', s.lower())
        for w in words + compounds:
            if w not in word_to_sentence:
                word_to_sentence[w] = s
    return word_to_sentence
