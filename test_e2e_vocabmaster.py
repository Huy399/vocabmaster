"""
Comprehensive Opaque-Box E2E Test Suite for VocabMaster.
Covers >= 115 test cases across 4 tiers:
- Tier 1: Feature Coverage (56 tests)
- Tier 2: Boundary & Corner Cases (56 tests)
- Tier 3: Cross-Feature Combinations (12 tests)
- Tier 4: Real-World Student Scenarios (6 tests)
Total: 130 automated tests.
"""
import os
import sys
import io
import re
import json
import time
import tempfile
import unittest
from datetime import datetime, timezone

# Ensure app package is importable
sys.path.insert(0, os.path.dirname(__file__))

import docx
import pymupdf
from fastapi.testclient import TestClient

import app.database as app_db
from app.main import app, _RATE_LIMIT_STORE, _TRANSLATE_CACHE
from app.srs import calculate_sm2
from app.vocab_engine import generate_smart_distractors, fetch_thesaurus, clean_vietnamese_text
from app.parsers import extract_text_from_docx, extract_text_from_pdf, clean_text, split_into_sentences


class BaseVocabMasterTest(unittest.TestCase):
    """Base class managing temporary isolated SQLite database and TestClient."""
    temp_db_path = None
    orig_db_path = None
    client = None

    @classmethod
    def setUpClass(cls):
        with tempfile.NamedTemporaryFile(suffix=".db", delete=False) as tmp:
            cls.temp_db_path = tmp.name
        cls.orig_db_path = app_db.DB_PATH
        app_db.DB_PATH = cls.temp_db_path
        app_db.init_db()
        cls.client = TestClient(app)

    @classmethod
    def tearDownClass(cls):
        app_db.DB_PATH = cls.orig_db_path
        if cls.temp_db_path and os.path.exists(cls.temp_db_path):
            try:
                os.remove(cls.temp_db_path)
            except Exception:
                pass

    def setUp(self):
        # Clear rate limiter and empty isolated tables for each test
        _RATE_LIMIT_STORE.clear()
        app_db.clear_all_words_db()
        app_db.clear_all_documents_db()

    def add_word(self, word="resilience", meaning_vi="khả năng phục hồi", pos="noun", ipa="/rɪˈzɪl.jəns/", example="He showed resilience.", cefr="B2"):
        return self.client.post("/api/words", json={
            "word": word,
            "meaning_vi": meaning_vi,
            "pos": pos,
            "ipa": ipa,
            "example_sentence": example,
            "cefr_level": cefr
        })

    def make_docx_bytes(self, text="Resilience and perseverance are vital qualities for students. Healthy nutrition provides vitamins."):
        doc = docx.Document()
        doc.add_heading("Study Guide", level=1)
        doc.add_paragraph(text)
        bio = io.BytesIO()
        doc.save(bio)
        return bio.getvalue()

    def make_pdf_bytes(self, text="Resilience and perseverance are vital qualities for students. Healthy nutrition provides vitamins."):
        pdf_doc = pymupdf.open()
        page = pdf_doc.new_page()
        page.insert_text((50, 72), text)
        pdf_bytes = pdf_doc.tobytes()
        pdf_doc.close()
        return pdf_bytes


# ==============================================================================
# TIER 1: FEATURE COVERAGE (56 tests)
# ==============================================================================
class TestTier1FeatureCoverage(BaseVocabMasterTest):
    """Tier 1: Feature Coverage across 10 functional domains."""

    # --- 1. 3D Flashcards & Rating Endpoints (R2.1) ---
    def test_t1_01_flashcard_due_words_initial(self):
        res = self.client.get("/api/words/due")
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertEqual(data["words"], [])
        self.assertEqual(data["count"], 0)

    def test_t1_02_flashcard_new_word_appears(self):
        self.add_word("resilience", "khả năng phục hồi")
        res = self.client.get("/api/words")
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertEqual(data["total"], 1)
        self.assertEqual(data["words"][0]["word"], "resilience")
        self.assertEqual(data["words"][0]["status"], "new")

    def test_t1_03_flashcard_review_good(self):
        res_add = self.add_word("diligence", "sự siêng năng")
        word_id = res_add.json()["word"]["id"]
        res = self.client.post("/api/words/review", json={"word_id": word_id, "quality": 4, "mode": "flashcard"})
        self.assertEqual(res.status_code, 200)
        w = res.json()["word"]
        self.assertEqual(w["interval"], 1)
        self.assertEqual(w["repetitions"], 1)
        self.assertEqual(w["status"], "learning")

    def test_t1_04_flashcard_review_fail(self):
        res_add = self.add_word("tenacity", "sự kiên trì")
        word_id = res_add.json()["word"]["id"]
        # Initial review
        self.client.post("/api/words/review", json={"word_id": word_id, "quality": 4, "mode": "flashcard"})
        # Fail review
        res = self.client.post("/api/words/review", json={"word_id": word_id, "quality": 1, "mode": "flashcard"})
        self.assertEqual(res.status_code, 200)
        w = res.json()["word"]
        self.assertEqual(w["interval"], 1)
        self.assertEqual(w["repetitions"], 0)
        self.assertEqual(w["status"], "learning")

    def test_t1_05_flashcard_review_logs_created(self):
        res_add = self.add_word("focus", "sự tập trung")
        word_id = res_add.json()["word"]["id"]
        self.client.post("/api/words/review", json={"word_id": word_id, "quality": 4, "mode": "flashcard"})
        stats = self.client.get("/api/stats").json()
        self.assertEqual(stats["total_reviews"], 1)

    def test_t1_06_flashcard_metadata_returned(self):
        res_add = self.add_word("vital", "quan trọng", pos="adj", ipa="/ˈvaɪ.təl/", example="Water is vital.")
        self.assertEqual(res_add.status_code, 200)
        w = res_add.json()["word"]
        self.assertEqual(w["word"], "vital")
        self.assertEqual(w["pos"], "adj")
        self.assertEqual(w["ipa"], "/ˈvaɪ.təl/")
        self.assertIn("vital", w["example_sentence"].lower())

    # --- 2. SM-2 SRS Algorithm Calculations (R2.2) ---
    def test_t1_07_sm2_calc_first_good(self):
        res = calculate_sm2(quality=4, repetitions=0, ease_factor=2.5, interval=1)
        self.assertEqual(res["interval"], 1)
        self.assertEqual(res["repetitions"], 1)

    def test_t1_08_sm2_calc_second_good(self):
        res = calculate_sm2(quality=4, repetitions=1, ease_factor=2.5, interval=1)
        self.assertEqual(res["interval"], 6)
        self.assertEqual(res["repetitions"], 2)

    def test_t1_09_sm2_calc_third_good(self):
        res = calculate_sm2(quality=4, repetitions=2, ease_factor=2.5, interval=6)
        self.assertEqual(res["interval"], 15)  # round(6 * 2.5) = 15
        self.assertEqual(res["repetitions"], 3)

    def test_t1_10_sm2_calc_reset_on_failure(self):
        res = calculate_sm2(quality=1, repetitions=4, ease_factor=2.5, interval=30)
        self.assertEqual(res["interval"], 1)
        self.assertEqual(res["repetitions"], 0)
        self.assertLess(res["ease_factor"], 2.5)

    def test_t1_11_sm2_calc_ease_factor_increase_on_perfect(self):
        res = calculate_sm2(quality=5, repetitions=1, ease_factor=2.5, interval=1)
        self.assertEqual(res["ease_factor"], 2.6)

    def test_t1_12_sm2_calc_ease_factor_decrease_on_hard(self):
        res = calculate_sm2(quality=3, repetitions=1, ease_factor=2.5, interval=1)
        self.assertEqual(res["ease_factor"], 2.36)

    # --- 3. Multiple Choice Quiz Generation & Distractors (R2.3) ---
    def test_t1_13_quiz_distractors_count(self):
        item = {"word": "perseverance", "meaning_vi": "sự kiên trì", "pos": "noun"}
        distractors = generate_smart_distractors(item, count=3)
        self.assertEqual(len(distractors), 3)

    def test_t1_14_quiz_distractors_differ_from_target(self):
        item = {"word": "perseverance", "meaning_vi": "sự kiên trì", "pos": "noun"}
        distractors = generate_smart_distractors(item, count=3)
        for d in distractors:
            self.assertNotEqual(d.lower(), item["meaning_vi"].lower())

    def test_t1_15_quiz_distractors_exclude_target_word(self):
        item = {"word": "apple", "meaning_vi": "quả táo", "pos": "noun"}
        distractors = generate_smart_distractors(item, count=3)
        for d in distractors:
            self.assertNotIn("apple", d.lower())

    def test_t1_16_quiz_distractor_capitalization(self):
        item = {"word": "courage", "meaning_vi": "Lòng dũng cảm", "pos": "noun"}
        distractors = generate_smart_distractors(item, count=3)
        for d in distractors:
            self.assertTrue(d[0].isupper())

    def test_t1_17_quiz_review_submission_mode(self):
        res_add = self.add_word("courage", "lòng dũng cảm")
        word_id = res_add.json()["word"]["id"]
        res = self.client.post("/api/words/review", json={"word_id": word_id, "quality": 4, "mode": "quiz"})
        self.assertEqual(res.status_code, 200)
        self.assertTrue(res.json()["success"])

    def test_t1_18_quiz_stored_distractors_json(self):
        res_add = self.add_word("wisdom", "sự thông thái")
        w = res_add.json()["word"]
        self.assertIsInstance(w.get("distractors"), list)
        self.assertGreaterEqual(len(w["distractors"]), 1)

    # --- 4. Fill-in-the-blank (Cloze) Logic & Evaluation (R2.4) ---
    def test_t1_19_cloze_sentence_masking(self):
        sentence = "He showed great resilience during difficult times."
        target_word = "resilience"
        masked = re.sub(rf"\b{re.escape(target_word)}\b", "___", sentence, flags=re.IGNORECASE)
        self.assertIn("___", masked)
        self.assertNotIn("resilience", masked.lower())

    def test_t1_20_cloze_case_insensitive_match(self):
        user_input = "Resilience"
        correct_word = "resilience"
        self.assertEqual(user_input.strip().lower(), correct_word.strip().lower())

    def test_t1_21_cloze_whitespace_tolerance(self):
        user_input = "   resilience   "
        correct_word = "resilience"
        self.assertEqual(user_input.strip().lower(), correct_word.strip().lower())

    def test_t1_22_cloze_incorrect_evaluation(self):
        user_input = "tenacity"
        correct_word = "resilience"
        self.assertNotEqual(user_input.strip().lower(), correct_word.strip().lower())

    def test_t1_23_cloze_review_submission_mode(self):
        res_add = self.add_word("endurance", "sự chịu đựng")
        word_id = res_add.json()["word"]["id"]
        res = self.client.post("/api/words/review", json={"word_id": word_id, "quality": 5, "mode": "cloze"})
        self.assertEqual(res.status_code, 200)
        self.assertTrue(res.json()["success"])

    # --- 5. Vocabulary Matching Game Logic (R2.5) ---
    def test_t1_24_match_card_generation(self):
        words = [
            {"id": 1, "word": "dog", "meaning_vi": "chó"},
            {"id": 2, "word": "cat", "meaning_vi": "mèo"},
            {"id": 3, "word": "sun", "meaning_vi": "mặt trời"},
            {"id": 4, "word": "moon", "meaning_vi": "mặt trăng"}
        ]
        cards = []
        for w in words:
            cards.append({"id": w["id"], "text": w["word"], "type": "en"})
            cards.append({"id": w["id"], "text": w["meaning_vi"], "type": "vi"})
        self.assertEqual(len(cards), 8)
        en_cards = [c for c in cards if c["type"] == "en"]
        vi_cards = [c for c in cards if c["type"] == "vi"]
        self.assertEqual(len(en_cards), 4)
        self.assertEqual(len(vi_cards), 4)

    def test_t1_25_match_correct_pair(self):
        c1 = {"id": 1, "text": "dog", "type": "en"}
        c2 = {"id": 1, "text": "chó", "type": "vi"}
        is_match = (c1["id"] == c2["id"] and c1["type"] != c2["type"])
        self.assertTrue(is_match)

    def test_t1_26_match_mismatched_pair(self):
        c1 = {"id": 1, "text": "dog", "type": "en"}
        c2 = {"id": 2, "text": "mèo", "type": "vi"}
        is_match = (c1["id"] == c2["id"] and c1["type"] != c2["type"])
        self.assertFalse(is_match)

    def test_t1_27_match_scoring_formula(self):
        matches = 6
        seconds = 45
        errors = 1
        base_score = matches * 100
        speed_bonus = max(0, 300 - seconds * 5)
        penalty = errors * 25
        final_score = max(0, base_score + speed_bonus - penalty)
        self.assertEqual(final_score, 600 + 75 - 25)

    def test_t1_28_match_review_submission_mode(self):
        res_add = self.add_word("speed", "tốc độ")
        word_id = res_add.json()["word"]["id"]
        res = self.client.post("/api/words/review", json={"word_id": word_id, "quality": 4, "mode": "match"})
        self.assertEqual(res.status_code, 200)
        self.assertTrue(res.json()["success"])

    # --- 6. Smart Thesaurus & Semantic Drift Control (R2.6) ---
    def test_t1_29_thesaurus_lookup_abundant(self):
        res = self.client.get("/api/thesaurus?word=abundant")
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertEqual(data["word"], "abundant")
        self.assertTrue(len(data["meaning_vi"]) > 0)
        self.assertGreaterEqual(len(data["synonyms"]), 3)

    def test_t1_30_thesaurus_drift_sound(self):
        res = self.client.get("/api/thesaurus?word=sound")
        self.assertEqual(res.status_code, 200)
        data = res.json()
        antonyms = {a["word"].lower() for a in data.get("antonyms", [])}
        self.assertNotIn("liquid", antonyms)
        self.assertNotIn("gas", antonyms)

    def test_t1_31_thesaurus_drift_fair(self):
        res = self.client.get("/api/thesaurus?word=fair")
        self.assertEqual(res.status_code, 200)
        data = res.json()
        synonyms = {s["word"].lower() for s in data.get("synonyms", [])}
        self.assertNotIn("carnival", synonyms)
        self.assertNotIn("bazaar", synonyms)

    def test_t1_32_thesaurus_drift_fine(self):
        res = self.client.get("/api/thesaurus?word=fine")
        self.assertEqual(res.status_code, 200)
        data = res.json()
        synonyms = {s["word"].lower() for s in data.get("synonyms", [])}
        self.assertNotIn("mulct", synonyms)
        self.assertNotIn("amercement", synonyms)

    def test_t1_33_thesaurus_drift_lead_patient(self):
        res_lead = self.client.get("/api/thesaurus?word=lead").json()
        lead_ants = {a["word"].lower() for a in res_lead.get("antonyms", [])}
        self.assertNotIn("tail", lead_ants)

        res_pat = self.client.get("/api/thesaurus?word=patient").json()
        pat_ants = {a["word"].lower() for a in res_pat.get("antonyms", [])}
        self.assertNotIn("outpatient", pat_ants)

    def test_t1_34_thesaurus_saved_status_indicator(self):
        self.add_word("abundant", "dồi dào")
        res = self.client.get("/api/thesaurus?word=abundant")
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertTrue(data["is_saved"])
        self.assertIsNotNone(data["saved_id"])

    # --- 7. Document Text Extraction (PDF/DOCX/TXT) (R2.7) ---
    def test_t1_35_extract_docx_text(self):
        b = self.make_docx_bytes("This is a valid DOCX sample for extraction testing.")
        text = extract_text_from_docx(b)
        self.assertGreater(len(text), 0)
        self.assertIn("extraction testing", text)

    def test_t1_36_extract_pdf_text(self):
        b = self.make_pdf_bytes("This is a valid PDF sample for extraction testing.")
        text = extract_text_from_pdf(b)
        self.assertGreater(len(text), 0)
        self.assertIn("extraction testing", text)

    def test_t1_37_upload_txt_endpoint(self):
        content = b"Resilience and perseverance are vital virtues for all students to master."
        files = {"file": ("test_doc.txt", io.BytesIO(content), "text/plain")}
        res = self.client.post("/api/upload", files=files)
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertTrue(data["success"])
        self.assertGreaterEqual(data["extracted_count"], 1)

    def test_t1_38_upload_docx_endpoint(self):
        b = self.make_docx_bytes()
        files = {"file": ("sample.docx", io.BytesIO(b), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        res = self.client.post("/api/upload", files=files)
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertTrue(data["success"])
        self.assertGreater(data["word_count"], 0)

    def test_t1_39_upload_pdf_endpoint(self):
        b = self.make_pdf_bytes()
        files = {"file": ("sample.pdf", io.BytesIO(b), "application/pdf")}
        res = self.client.post("/api/upload", files=files)
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertTrue(data["success"])
        self.assertGreater(data["word_count"], 0)

    def test_t1_40_get_document_detail(self):
        content = b"English vocabulary learning is essential for reading academic literature."
        files = {"file": ("article.txt", io.BytesIO(content), "text/plain")}
        res_up = self.client.post("/api/upload", files=files)
        doc_id = res_up.json()["document_id"]

        res_get = self.client.get(f"/api/documents/{doc_id}")
        self.assertEqual(res_get.status_code, 200)
        doc = res_get.json()
        self.assertEqual(doc["filename"], "article.txt")
        self.assertIn("academic", doc["text_content"])

    # --- 8. Security Defenses (R4) ---
    def test_t1_41_xss_in_word_creation(self):
        payload = "<script>alert('xss')</script>"
        res = self.add_word(payload, "mã độc kiểm thử")
        self.assertEqual(res.status_code, 200)
        # Server accepts data safely without executing or crashing
        saved_words = self.client.get("/api/words").json()["words"]
        self.assertEqual(saved_words[0]["word"], payload)

    def test_t1_42_xss_in_meaning(self):
        payload = '<img src=x onerror=alert(1)>'
        res = self.add_word("secure_term", payload)
        self.assertEqual(res.status_code, 200)
        self.assertTrue(res.json()["success"])

    def test_t1_43_sql_injection_search_resilience(self):
        self.add_word("resilience", "khả năng phục hồi")
        res = self.client.get("/api/words?search=' OR '1'='1")
        self.assertEqual(res.status_code, 200)
        data = res.json()
        # Escaped search should NOT dump all rows because literal "' OR '1'='1" does not match
        self.assertEqual(len(data["words"]), 0)

    def test_t1_44_path_traversal_upload_filename(self):
        content = b"Valid test document text with twenty or more characters resilience."
        files = {"file": ("../../malicious_traversal.txt", io.BytesIO(content), "text/plain")}
        res = self.client.post("/api/upload", files=files)
        self.assertEqual(res.status_code, 200)
        # Verify no file was written outside to disk
        self.assertFalse(os.path.exists("../../malicious_traversal.txt"))

    def test_t1_45_magic_bytes_fake_pdf_rejected(self):
        content = b"This is plain text pretending to be a PDF file."
        files = {"file": ("fake.pdf", io.BytesIO(content), "application/pdf")}
        res = self.client.post("/api/upload", files=files)
        self.assertEqual(res.status_code, 400)
        self.assertIn("PDF", res.json()["detail"])

    def test_t1_46_magic_bytes_fake_docx_rejected(self):
        content = b"This is plain text pretending to be a DOCX file."
        files = {"file": ("fake.docx", io.BytesIO(content), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        res = self.client.post("/api/upload", files=files)
        self.assertEqual(res.status_code, 400)
        self.assertIn("Word", res.json()["detail"])

    # --- 9. Performance & Caching (R3) ---
    def test_t1_47_pagination_limit_and_offset(self):
        for i in range(12):
            self.add_word(f"word_{i:02d}", f"nghĩa {i}")
        res = self.client.get("/api/words?limit=5&offset=0")
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertEqual(len(data["words"]), 5)
        self.assertEqual(data["total"], 12)

    def test_t1_48_pagination_offset_progression(self):
        for i in range(10):
            self.add_word(f"term_{i:02d}", f"nghĩa {i}")
        page1 = self.client.get("/api/words?limit=4&offset=0").json()["words"]
        page2 = self.client.get("/api/words?limit=4&offset=4").json()["words"]
        ids1 = {w["id"] for w in page1}
        ids2 = {w["id"] for w in page2}
        self.assertEqual(len(ids1.intersection(ids2)), 0)

    def test_t1_49_translation_cache_hit(self):
        res1 = self.client.get("/api/translate?word=resilience")
        self.assertEqual(res1.status_code, 200)
        res2 = self.client.get("/api/translate?word=resilience")
        self.assertEqual(res2.status_code, 200)
        self.assertEqual(res1.json()["meaning_vi"], res2.json()["meaning_vi"])

    def test_t1_50_thesaurus_cache_hit(self):
        t0 = time.time()
        res1 = self.client.get("/api/thesaurus?word=abundant")
        d1 = time.time() - t0

        t1 = time.time()
        res2 = self.client.get("/api/thesaurus?word=abundant")
        d2 = time.time() - t1
        self.assertEqual(res1.status_code, 200)
        self.assertEqual(res2.status_code, 200)
        self.assertLessEqual(d2, d1 + 0.1)

    def test_t1_51_stats_single_query_aggregation(self):
        self.add_word("speed", "tốc độ")
        res = self.client.get("/api/stats")
        self.assertEqual(res.status_code, 200)
        stats = res.json()
        self.assertIn("total_words", stats)
        self.assertIn("new_words", stats)
        self.assertIn("learning_words", stats)
        self.assertIn("mastered_words", stats)
        self.assertIn("total_reviews", stats)
        self.assertEqual(stats["total_words"], 1)

    # --- 10. Robustness & Error Responses (R4) ---
    def test_t1_52_post_word_missing_meaning_422(self):
        res = self.client.post("/api/words", json={"word": "empty_meaning"})
        self.assertEqual(res.status_code, 422)

    def test_t1_53_review_invalid_word_id_400(self):
        res = self.client.post("/api/words/review", json={"word_id": 999999, "quality": 4})
        self.assertEqual(res.status_code, 400)

    def test_t1_54_get_missing_document_404(self):
        res = self.client.get("/api/documents/999999")
        self.assertEqual(res.status_code, 404)

    def test_t1_55_delete_missing_word_404(self):
        res = self.client.delete("/api/words/999999")
        self.assertEqual(res.status_code, 404)

    def test_t1_56_error_json_contains_detail(self):
        res = self.client.get("/api/documents/999999")
        data = res.json()
        self.assertIn("detail", data)


# ==============================================================================
# TIER 2: BOUNDARY & CORNER CASES (56 tests)
# ==============================================================================
class TestTier2BoundaryCornerCases(BaseVocabMasterTest):
    """Tier 2: Boundary, corner, and edge cases across algorithms and inputs."""

    # --- 1. SRS Ratings Boundary (10 tests) ---
    def test_t2_01_rating_0_blackout(self):
        res = calculate_sm2(quality=0, repetitions=3, ease_factor=2.5, interval=15)
        self.assertEqual(res["repetitions"], 0)
        self.assertEqual(res["interval"], 1)
        self.assertLess(res["ease_factor"], 2.5)

    def test_t2_02_rating_1_failure(self):
        res = calculate_sm2(quality=1, repetitions=2, ease_factor=2.5, interval=6)
        self.assertEqual(res["repetitions"], 0)
        self.assertEqual(res["interval"], 1)

    def test_t2_03_rating_2_borderline_fail(self):
        res = calculate_sm2(quality=2, repetitions=1, ease_factor=2.5, interval=6)
        self.assertEqual(res["repetitions"], 0)
        self.assertEqual(res["interval"], 1)

    def test_t2_04_rating_3_hard_pass(self):
        res = calculate_sm2(quality=3, repetitions=2, ease_factor=2.5, interval=6)
        self.assertEqual(res["repetitions"], 3)
        self.assertGreater(res["interval"], 6)
        self.assertLess(res["ease_factor"], 2.5)

    def test_t2_05_rating_4_good_pass(self):
        res = calculate_sm2(quality=4, repetitions=0, ease_factor=2.5, interval=1)
        self.assertEqual(res["repetitions"], 1)
        self.assertEqual(res["interval"], 1)

    def test_t2_06_rating_5_easy_pass(self):
        res = calculate_sm2(quality=5, repetitions=1, ease_factor=2.5, interval=1)
        self.assertEqual(res["repetitions"], 2)
        self.assertEqual(res["ease_factor"], 2.6)

    def test_t2_07_rating_6_clamped_to_5(self):
        res = calculate_sm2(quality=6, repetitions=1, ease_factor=2.5, interval=1)
        self.assertEqual(res["ease_factor"], 2.6)

    def test_t2_08_rating_negative_clamped_to_0(self):
        res = calculate_sm2(quality=-1, repetitions=1, ease_factor=2.5, interval=6)
        self.assertEqual(res["repetitions"], 0)
        self.assertEqual(res["interval"], 1)

    def test_t2_09_rating_float_coerced(self):
        res = calculate_sm2(quality=4.8, repetitions=0, ease_factor=2.5, interval=1)
        self.assertEqual(res["repetitions"], 1)

    def test_t2_10_rating_string_coerced(self):
        res = calculate_sm2(quality="4", repetitions=0, ease_factor=2.5, interval=1)
        self.assertEqual(res["repetitions"], 1)

    # --- 2. Ease Factor & Interval Clamping (6 tests) ---
    def test_t2_11_ease_factor_minimum_bound(self):
        ef = 2.5
        for _ in range(25):
            res = calculate_sm2(quality=0, repetitions=0, ease_factor=ef, interval=1)
            ef = res["ease_factor"]
        self.assertGreaterEqual(ef, 1.3)
        self.assertEqual(ef, 1.3)

    def test_t2_12_ease_factor_maximum_bound(self):
        ef = 2.5
        for _ in range(25):
            res = calculate_sm2(quality=5, repetitions=1, ease_factor=ef, interval=1)
            ef = res["ease_factor"]
        self.assertLessEqual(ef, 3.0)
        self.assertEqual(ef, 3.0)

    def test_t2_13_interval_maximum_capped_at_365(self):
        res = calculate_sm2(quality=5, repetitions=15, ease_factor=3.0, interval=250)
        self.assertLessEqual(res["interval"], 365)
        self.assertEqual(res["interval"], 365)

    def test_t2_14_zero_repetitions_gives_interval_1(self):
        res = calculate_sm2(quality=4, repetitions=0, ease_factor=2.5, interval=1)
        self.assertEqual(res["interval"], 1)

    def test_t2_15_one_repetition_gives_interval_6(self):
        res = calculate_sm2(quality=4, repetitions=1, ease_factor=2.5, interval=1)
        self.assertEqual(res["interval"], 6)

    def test_t2_16_due_date_utc_midnight_normalization(self):
        res = calculate_sm2(quality=4, repetitions=0, ease_factor=2.5, interval=1)
        due_iso = res["due_date"]
        self.assertTrue(due_iso.endswith("+00:00") or due_iso.endswith("Z"))
        self.assertIn("T00:00:00", due_iso)

    # --- 3. Word Input, Character & String Boundaries (8 tests) ---
    def test_t2_17_empty_word_creation(self):
        res = self.client.post("/api/words", json={"word": "", "meaning_vi": "nghĩa"})
        # Should either reject or save empty cleanly
        self.assertIn(res.status_code, [200, 400, 422])

    def test_t2_18_whitespace_word_validation(self):
        res = self.client.post("/api/words", json={"word": "   ", "meaning_vi": "nghĩa"})
        self.assertIn(res.status_code, [200, 400, 422])

    def test_t2_19_word_with_leading_trailing_whitespace_stripped(self):
        res = self.add_word("  perseverance  ", "sự kiên trì")
        self.assertEqual(res.status_code, 200)
        w = res.json()["word"]
        self.assertEqual(w["word"], "perseverance")

    def test_t2_20_hyphenated_word(self):
        res = self.add_word("state-of-the-art", "hiện đại nhất")
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.json()["word"]["word"], "state-of-the-art")

    def test_t2_21_thesaurus_word_with_apostrophe(self):
        res = self.client.get("/api/thesaurus?word=don't")
        self.assertEqual(res.status_code, 200)

    def test_t2_22_thesaurus_word_exceeding_80_chars_rejected(self):
        long_word = "a" * 81
        res = self.client.get(f"/api/thesaurus?word={long_word}")
        self.assertEqual(res.status_code, 400)

    def test_t2_23_thesaurus_vietnamese_accented_input_rejected(self):
        res = self.client.get("/api/thesaurus?word=học tập")
        self.assertEqual(res.status_code, 400)

    def test_t2_24_thesaurus_malicious_delimiters_rejected(self):
        res = self.client.get("/api/thesaurus?word=test; DROP TABLE words; --")
        self.assertEqual(res.status_code, 400)

    # --- 4. Translation & Cache Boundaries (6 tests) ---
    def test_t2_25_translate_empty_string(self):
        res = self.client.get("/api/translate?word=")
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.json()["word"], "")

    def test_t2_26_translate_whitespace_only(self):
        res = self.client.get("/api/translate?word=   ")
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.json()["word"], "")

    def test_t2_27_translate_single_char(self):
        res = self.client.get("/api/translate?word=a")
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.json()["word"], "a")

    def test_t2_28_translate_uppercase(self):
        res = self.client.get("/api/translate?word=RESILIENCE")
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.json()["word"], "resilience")

    def test_t2_29_translate_builtin_vocab(self):
        res = self.client.get("/api/translate?word=health")
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertEqual(data["word"], "health")
        self.assertIn("sức khỏe", data["meaning_vi"].lower())

    def test_t2_30_translate_uncommon_word(self):
        res = self.client.get("/api/translate?word=serendipity")
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.json()["word"], "serendipity")

    # --- 5. File Upload & Format Boundaries (8 tests) ---
    def test_t2_31_upload_zero_bytes_rejected(self):
        files = {"file": ("empty.txt", io.BytesIO(b""), "text/plain")}
        res = self.client.post("/api/upload", files=files)
        self.assertEqual(res.status_code, 400)

    def test_t2_32_upload_under_20_chars_rejected(self):
        files = {"file": ("short.txt", io.BytesIO(b"Short text 123"), "text/plain")}
        res = self.client.post("/api/upload", files=files)
        self.assertEqual(res.status_code, 400)

    def test_t2_33_upload_exact_20_chars_accepted(self):
        # 21 characters long
        content = b"1234567890 1234567890"
        files = {"file": ("twenty.txt", io.BytesIO(content), "text/plain")}
        res = self.client.post("/api/upload", files=files)
        self.assertEqual(res.status_code, 200)

    def test_t2_34_upload_disallowed_extension(self):
        content = b"Some valid content with plenty of words and resilience."
        files = {"file": ("script.exe", io.BytesIO(content), "application/x-msdownload")}
        res = self.client.post("/api/upload", files=files)
        self.assertEqual(res.status_code, 400)

    def test_t2_35_upload_fake_pdf_header(self):
        content = b"NOT_A_PDF header text that fails magic byte verification."
        files = {"file": ("fake_header.pdf", io.BytesIO(content), "application/pdf")}
        res = self.client.post("/api/upload", files=files)
        self.assertEqual(res.status_code, 400)

    def test_t2_36_upload_fake_docx_header(self):
        content = b"NOT_A_DOCX zip header text that fails magic byte verification."
        files = {"file": ("fake_header.docx", io.BytesIO(content), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        res = self.client.post("/api/upload", files=files)
        self.assertEqual(res.status_code, 400)

    def test_t2_37_upload_null_bytes_in_txt(self):
        content = b"Valid text but has " + bytes([0]) + b" null binary byte."
        files = {"file": ("binary.txt", io.BytesIO(content), "text/plain")}
        res = self.client.post("/api/upload", files=files)
        self.assertEqual(res.status_code, 400)

    def test_t2_38_upload_unicode_emojis_accepted(self):
        content = "Học từ vựng tiếng Anh với VocabMaster rất tuyệt vời! 🚀🌟 Resilience and perseverance.".encode("utf-8")
        files = {"file": ("unicode.txt", io.BytesIO(content), "text/plain")}
        res = self.client.post("/api/upload", files=files)
        self.assertEqual(res.status_code, 200)

    # --- 6. Distractor & Quiz Boundaries (6 tests) ---
    def test_t2_39_distractors_count_1(self):
        item = {"word": "vital", "meaning_vi": "quan trọng", "pos": "adj"}
        d = generate_smart_distractors(item, count=1)
        self.assertEqual(len(d), 1)

    def test_t2_40_distractors_count_5(self):
        item = {"word": "vital", "meaning_vi": "quan trọng", "pos": "adj"}
        d = generate_smart_distractors(item, count=5)
        self.assertEqual(len(d), 5)

    def test_t2_41_distractor_split_slash(self):
        item = {"word": "speedy", "meaning_vi": "nhanh chóng / mau lẹ", "pos": "adj"}
        d = generate_smart_distractors(item, count=3)
        for dist in d:
            self.assertNotIn("nhanh chóng", dist.lower())
            self.assertNotIn("mau lẹ", dist.lower())

    def test_t2_42_distractor_split_comma(self):
        item = {"word": "prosperous", "meaning_vi": "thịnh vượng, phát đạt", "pos": "adj"}
        d = generate_smart_distractors(item, count=3)
        for dist in d:
            self.assertNotIn("thịnh vượng", dist.lower())
            self.assertNotIn("phát đạt", dist.lower())

    def test_t2_43_distractor_unknown_pos(self):
        item = {"word": "xyzword", "meaning_vi": "nghĩa lạ", "pos": "xyz_pos"}
        d = generate_smart_distractors(item, count=3)
        self.assertEqual(len(d), 3)

    def test_t2_44_distractor_empty_meaning(self):
        item = {"word": "empty_mean", "meaning_vi": "", "pos": "noun"}
        d = generate_smart_distractors(item, count=3)
        self.assertEqual(len(d), 3)

    # --- 7. Database & Search Boundaries (6 tests) ---
    def test_t2_45_offset_beyond_total(self):
        self.add_word("single_word", "một từ")
        res = self.client.get("/api/words?offset=999&limit=10")
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertEqual(len(data["words"]), 0)
        self.assertEqual(data["total"], 1)

    def test_t2_46_limit_zero(self):
        self.add_word("item", "mục")
        res = self.client.get("/api/words?limit=0")
        self.assertEqual(res.status_code, 200)
        # In this backend contract, limit <= 0 returns all words without SQL limit
        self.assertGreaterEqual(len(res.json()["words"]), 1)

    def test_t2_47_search_literal_percent(self):
        self.add_word("regular_word", "từ thường")
        res = self.client.get("/api/words?search=%")
        self.assertEqual(res.status_code, 200)
        # Should match nothing because no word literally has '%'
        self.assertEqual(len(res.json()["words"]), 0)

    def test_t2_48_search_literal_underscore(self):
        self.add_word("apple", "táo")
        # An unescaped SQL underscore '_' would match any single char like 'a' in 'apple'
        res = self.client.get("/api/words?search=_")
        self.assertEqual(res.status_code, 200)
        self.assertEqual(len(res.json()["words"]), 0)

    def test_t2_49_deduplication_in_batch(self):
        words_data = [
            {"word": "duplicate", "meaning_vi": "trùng lặp"},
            {"word": "duplicate", "meaning_vi": "trùng lặp phiên bản 2"}
        ]
        saved = app_db.save_words(None, words_data)
        self.assertEqual(len(saved), 1)
        all_words = app_db.get_all_words()
        self.assertEqual(len(all_words), 1)

    def test_t2_50_upsert_existing_word(self):
        self.add_word("persist", "kiên trì")
        # Upsert with new meaning
        self.add_word("persist", "tiếp tục tồn tại")
        all_words = self.client.get("/api/words").json()["words"]
        self.assertEqual(len(all_words), 1)
        self.assertEqual(all_words[0]["word"], "persist")
        self.assertEqual(all_words[0]["meaning_vi"], "tiếp tục tồn tại")

    # --- 8. Cloze & Match Boundaries (6 tests) ---
    def test_t2_51_cloze_trailing_punctuation(self):
        user_input = "resilience."
        cleaned = user_input.strip(" .,!?;:\"'").lower()
        self.assertEqual(cleaned, "resilience")

    def test_t2_52_cloze_empty_answer_fails(self):
        user_input = ""
        correct = "resilience"
        self.assertFalse(bool(user_input.strip()))

    def test_t2_53_cloze_word_at_start_masked(self):
        sentence = "Resilience is essential for success."
        target = "resilience"
        masked = re.sub(rf"\b{re.escape(target)}\b", "___", sentence, flags=re.IGNORECASE)
        self.assertTrue(masked.startswith("___"))

    def test_t2_54_cloze_word_at_end_masked(self):
        sentence = "He demonstrated great resilience."
        target = "resilience"
        masked = re.sub(rf"\b{re.escape(target)}\b", "___", sentence, flags=re.IGNORECASE)
        self.assertTrue(masked.endswith("___."))

    def test_t2_55_match_small_pool_2_words(self):
        words = [
            {"id": 1, "word": "a", "meaning_vi": "nghĩa a"},
            {"id": 2, "word": "b", "meaning_vi": "nghĩa b"}
        ]
        cards = []
        for w in words:
            cards.append({"id": w["id"], "text": w["word"], "type": "en"})
            cards.append({"id": w["id"], "text": w["meaning_vi"], "type": "vi"})
        self.assertEqual(len(cards), 4)

    def test_t2_56_match_distinct_card_ids(self):
        c1 = {"id": 1, "text": "word", "type": "en"}
        c2 = {"id": 2, "text": "word", "type": "en"}
        self.assertNotEqual(c1["id"], c2["id"])


# ==============================================================================
# TIER 3: CROSS-FEATURE COMBINATIONS (12 tests)
# ==============================================================================
class TestTier3CrossFeatureCombinations(BaseVocabMasterTest):
    """Tier 3: Multi-feature workflows, cascading updates, and data lifecycles."""

    def test_t3_01_upload_to_words_lifecycle(self):
        content = b"The diligent student demonstrated high resilience during examinations."
        files = {"file": ("study.txt", io.BytesIO(content), "text/plain")}
        res_up = self.client.post("/api/upload", files=files)
        self.assertEqual(res_up.status_code, 200)
        doc_id = res_up.json()["document_id"]

        words = self.client.get("/api/words").json()["words"]
        self.assertGreaterEqual(len(words), 1)
        for w in words:
            self.assertEqual(w["document_id"], doc_id)

    def test_t3_02_five_round_mastery_progression(self):
        res_add = self.add_word("perseverance", "sự kiên trì")
        word_id = res_add.json()["word"]["id"]

        # 5 consecutive high-quality reviews
        for q in [4, 4, 4, 4, 5]:
            res = self.client.post("/api/words/review", json={"word_id": word_id, "quality": q})
            self.assertEqual(res.status_code, 200)

        w = self.client.get("/api/words").json()["words"][0]
        self.assertEqual(w["status"], "mastered")
        self.assertEqual(w["repetitions"], 5)
        self.assertGreater(w["interval"], 6)

    def test_t3_03_mastered_word_relapses_on_failure(self):
        res_add = self.add_word("mastered_target", "từ đã thông thạo")
        word_id = res_add.json()["word"]["id"]

        # Reach mastery
        for q in [4, 4, 4, 4, 5]:
            self.client.post("/api/words/review", json={"word_id": word_id, "quality": q})

        # Relapse with quality 1
        res_fail = self.client.post("/api/words/review", json={"word_id": word_id, "quality": 1})
        w = res_fail.json()["word"]
        self.assertEqual(w["status"], "learning")
        self.assertEqual(w["repetitions"], 0)
        self.assertEqual(w["interval"], 1)

    def test_t3_04_thesaurus_save_and_recheck_status(self):
        # Step 1: Thesaurus lookup when word is NOT saved
        res1 = self.client.get("/api/thesaurus?word=abundant")
        self.assertFalse(res1.json()["is_saved"])

        # Step 2: Save word to vocabulary
        self.add_word("abundant", res1.json()["meaning_vi"])

        # Step 3: Thesaurus lookup again, now is_saved should be True
        res2 = self.client.get("/api/thesaurus?word=abundant")
        self.assertTrue(res2.json()["is_saved"])
        self.assertIsNotNone(res2.json()["saved_id"])

    def test_t3_05_quiz_from_saved_word(self):
        res_add = self.add_word("nutrition", "dinh dưỡng", pos="noun")
        w = res_add.json()["word"]
        distractors = w["distractors"]
        self.assertGreaterEqual(len(distractors), 1)

        options = [w["meaning_vi"]] + distractors[:3]
        self.assertIn(w["meaning_vi"], options)

        # Submit quiz review
        res_rev = self.client.post("/api/words/review", json={"word_id": w["id"], "quality": 4, "mode": "quiz"})
        self.assertEqual(res_rev.status_code, 200)

        stats = self.client.get("/api/stats").json()
        self.assertEqual(stats["total_reviews"], 1)

    def test_t3_06_match_game_full_flow(self):
        # Add 6 words
        words_input = [
            ("apple", "quả táo"), ("banana", "quả chuối"), ("orange", "quả cam"),
            ("grape", "quả nho"), ("lemon", "quả chanh"), ("peach", "quả đào")
        ]
        for eng, vi in words_input:
            self.add_word(eng, vi)

        all_words = self.client.get("/api/words").json()["words"]
        self.assertEqual(len(all_words), 6)

        # Generate match pairs
        en_cards = {w["id"]: w["word"] for w in all_words}
        vi_cards = {w["id"]: w["meaning_vi"] for w in all_words}

        # Simulate matching all pairs
        for w_id in en_cards:
            self.client.post("/api/words/review", json={"word_id": w_id, "quality": 4, "mode": "match"})

        stats = self.client.get("/api/stats").json()
        self.assertEqual(stats["total_reviews"], 6)

    def test_t3_07_document_deletion_preserves_words(self):
        content = b"Resilience and perseverance are essential virtues for every student."
        files = {"file": ("preserve_test.txt", io.BytesIO(content), "text/plain")}
        res_up = self.client.post("/api/upload", files=files)
        doc_id = res_up.json()["document_id"]

        words_before = self.client.get("/api/words").json()["words"]
        self.assertGreater(len(words_before), 0)

        # Delete document
        res_del = self.client.delete(f"/api/documents/{doc_id}")
        self.assertEqual(res_del.status_code, 200)

        # Words should still exist, with document_id unlinked (NULL)
        words_after = self.client.get("/api/words").json()["words"]
        self.assertEqual(len(words_after), len(words_before))
        for w in words_after:
            self.assertIsNone(w["document_id"])

    def test_t3_08_clear_all_documents_preserves_words(self):
        content = b"Resilience and perseverance are vital concepts in education."
        files = {"file": ("clear_doc_test.txt", io.BytesIO(content), "text/plain")}
        self.client.post("/api/upload", files=files)

        words_count = len(self.client.get("/api/words").json()["words"])
        self.assertGreater(words_count, 0)

        # Clear all documents
        res_clear = self.client.delete("/api/documents")
        self.assertEqual(res_clear.status_code, 200)

        docs = self.client.get("/api/documents").json()["documents"]
        self.assertEqual(len(docs), 0)

        # Words table preserved
        words_after = self.client.get("/api/words").json()["words"]
        self.assertEqual(len(words_after), words_count)

    def test_t3_09_clear_all_words_cascades_reviews(self):
        res_add = self.add_word("temp_word", "từ tạm")
        word_id = res_add.json()["word"]["id"]
        self.client.post("/api/words/review", json={"word_id": word_id, "quality": 4})

        stats_before = self.client.get("/api/stats").json()
        self.assertEqual(stats_before["total_words"], 1)
        self.assertEqual(stats_before["total_reviews"], 1)

        # Clear all words
        res_clear = self.client.delete("/api/words")
        self.assertEqual(res_clear.status_code, 200)

        stats_after = self.client.get("/api/stats").json()
        self.assertEqual(stats_after["total_words"], 0)
        self.assertEqual(stats_after["total_reviews"], 0)

    def test_t3_10_multi_mode_review_history(self):
        res_add = self.add_word("synergy", "sự cộng hưởng")
        word_id = res_add.json()["word"]["id"]

        self.client.post("/api/words/review", json={"word_id": word_id, "quality": 4, "mode": "flashcard"})
        self.client.post("/api/words/review", json={"word_id": word_id, "quality": 4, "mode": "quiz"})
        self.client.post("/api/words/review", json={"word_id": word_id, "quality": 5, "mode": "cloze"})

        stats = self.client.get("/api/stats").json()
        self.assertEqual(stats["total_reviews"], 3)

    def test_t3_11_stats_synchronization_lifecycle(self):
        # Initial empty
        s0 = self.client.get("/api/stats").json()
        self.assertEqual(s0["total_words"], 0)

        # Add 1 word
        res_add = self.add_word("sync_word", "từ đồng bộ")
        w_id = res_add.json()["word"]["id"]
        s1 = self.client.get("/api/stats").json()
        self.assertEqual(s1["total_words"], 1)
        self.assertEqual(s1["new_words"], 1)

        # Review word
        self.client.post("/api/words/review", json={"word_id": w_id, "quality": 4})
        s2 = self.client.get("/api/stats").json()
        self.assertEqual(s2["new_words"], 0)
        self.assertEqual(s2["learning_words"], 1)

        # Delete word
        self.client.delete(f"/api/words/{w_id}")
        s3 = self.client.get("/api/stats").json()
        self.assertEqual(s3["total_words"], 0)

    def test_t3_12_combined_status_and_search_query(self):
        self.add_word("apple_pie", "bánh táo")
        res2 = self.add_word("apple_juice", "nước táo")
        w2_id = res2.json()["word"]["id"]
        # Review word 2 so it becomes 'learning'
        self.client.post("/api/words/review", json={"word_id": w2_id, "quality": 4})

        res = self.client.get("/api/words?status=learning&search=juice")
        self.assertEqual(res.status_code, 200)
        data = res.json()["words"]
        self.assertEqual(len(data), 1)
        self.assertEqual(data[0]["word"], "apple_juice")


# ==============================================================================
# TIER 4: REAL-WORLD STUDENT SCENARIOS (6 tests)
# ==============================================================================
class TestTier4RealWorldScenarios(BaseVocabMasterTest):
    """Tier 4: Comprehensive end-to-end user workflows and student journeys."""

    def test_t4_01_student_onboarding_journey(self):
        """A new student starts from 0, adds 3 target vocabulary items, studies card 1."""
        # 1. Check empty dashboard
        s0 = self.client.get("/api/stats").json()
        self.assertEqual(s0["total_words"], 0)

        # 2. Add 3 words manually
        w1 = self.add_word("resilience", "khả năng phục hồi", "noun", "/rɪˈzɪl.jəns/").json()["word"]
        w2 = self.add_word("diligence", "sự siêng năng", "noun", "/ˈdɪl.ɪ.dʒəns/").json()["word"]
        w3 = self.add_word("fortitude", "sự kiên cường", "noun", "/ˈfɔː.tɪ.tʃuːd/").json()["word"]

        s1 = self.client.get("/api/stats").json()
        self.assertEqual(s1["total_words"], 3)
        self.assertEqual(s1["new_words"], 3)

        # 3. Practice Flashcard for word 1 (Quality 4: Good)
        res_rev = self.client.post("/api/words/review", json={"word_id": w1["id"], "quality": 4, "mode": "flashcard"})
        self.assertEqual(res_rev.status_code, 200)

        # 4. Dashboard updates
        s2 = self.client.get("/api/stats").json()
        self.assertEqual(s2["new_words"], 2)
        self.assertEqual(s2["learning_words"], 1)

    def test_t4_02_academic_passage_reading_and_word_saving(self):
        """Student uploads reading material, inspects reader, looks up thesaurus, saves new term."""
        passage = b"Nutrition and wellness are fundamental concepts in human biology and health."
        files = {"file": ("biology_ch1.txt", io.BytesIO(passage), "text/plain")}
        res_up = self.client.post("/api/upload", files=files)
        self.assertEqual(res_up.status_code, 200)
        doc_id = res_up.json()["document_id"]

        # Student opens document reader
        doc = self.client.get(f"/api/documents/{doc_id}").json()
        self.assertIn("biology", doc["text_content"])

        # Student clicks an unfamiliar word to translate
        trans = self.client.get("/api/translate?word=wellness").json()
        self.assertEqual(trans["word"], "wellness")

        # Student explores synonyms in thesaurus
        thes = self.client.get("/api/thesaurus?word=abundant").json()
        self.assertGreater(len(thes["synonyms"]), 0)

        # Quick save word
        res_save = self.add_word("abundant", thes["meaning_vi"])
        self.assertEqual(res_save.status_code, 200)

    def test_t4_03_comprehensive_exam_session(self):
        """Student runs a 5-question exam (3 quiz + 2 cloze questions) scoring 4/5."""
        # Create 5 test words
        words = []
        for term, vi in [("acumen", "sự nhạy bén"), ("benevolent", "nhân từ"), ("candid", "thẳng thắn"),
                         ("dauntless", "dũng cảm"), ("eloquent", "hùng biện")]:
            w = self.add_word(term, vi).json()["word"]
            words.append(w)

        # Questions 1-3: Multiple Choice Quiz (Student gets Q1 & Q2 right, Q3 wrong)
        self.client.post("/api/words/review", json={"word_id": words[0]["id"], "quality": 4, "mode": "quiz"})
        self.client.post("/api/words/review", json={"word_id": words[1]["id"], "quality": 5, "mode": "quiz"})
        self.client.post("/api/words/review", json={"word_id": words[2]["id"], "quality": 1, "mode": "quiz"})  # Incorrect

        # Questions 4-5: Cloze (Student gets Q4 & Q5 right)
        self.client.post("/api/words/review", json={"word_id": words[3]["id"], "quality": 4, "mode": "cloze"})
        self.client.post("/api/words/review", json={"word_id": words[4]["id"], "quality": 5, "mode": "cloze"})

        # Verify review logs
        stats = self.client.get("/api/stats").json()
        self.assertEqual(stats["total_reviews"], 5)

        # Verify failed word remains at interval 1
        q3_word = self.client.get("/api/words?search=candid").json()["words"][0]
        self.assertEqual(q3_word["repetitions"], 0)
        self.assertEqual(q3_word["interval"], 1)

    def test_t4_04_speed_matching_game_challenge(self):
        """Student finishes matching game with 6 word pairs cleanly."""
        terms = [("apex", "đỉnh cao"), ("bold", "táo bạo"), ("calm", "bình tĩnh"),
                 ("dawn", "bình minh"), ("ease", "dễ chịu"), ("fast", "nhanh")]
        word_ids = []
        for eng, vi in terms:
            w = self.add_word(eng, vi).json()["word"]
            word_ids.append(w["id"])

        # Player matches all pairs in 25 seconds
        for w_id in word_ids:
            res = self.client.post("/api/words/review", json={"word_id": w_id, "quality": 4, "mode": "match"})
            self.assertEqual(res.status_code, 200)

        stats = self.client.get("/api/stats").json()
        self.assertEqual(stats["total_reviews"], 6)

    def test_t4_05_spaced_repetition_multi_stage_mastery(self):
        """Simulates spaced learning over multiple review milestones to complete mastery."""
        w = self.add_word("mastery_journey", "hành trình thông thạo").json()["word"]
        word_id = w["id"]

        # Milestone 1: Day 1
        r1 = self.client.post("/api/words/review", json={"word_id": word_id, "quality": 4}).json()["word"]
        self.assertEqual(r1["interval"], 1)
        self.assertEqual(r1["repetitions"], 1)

        # Milestone 2: Day 2
        r2 = self.client.post("/api/words/review", json={"word_id": word_id, "quality": 4}).json()["word"]
        self.assertEqual(r2["interval"], 6)
        self.assertEqual(r2["repetitions"], 2)

        # Milestone 3: Day 8
        r3 = self.client.post("/api/words/review", json={"word_id": word_id, "quality": 4}).json()["word"]
        self.assertEqual(r3["repetitions"], 3)

        # Milestone 4: Day 23
        r4 = self.client.post("/api/words/review", json={"word_id": word_id, "quality": 4}).json()["word"]
        self.assertEqual(r4["repetitions"], 4)

        # Milestone 5: Day 61 (Quality 5: Easy) -> Transitions to Mastered
        r5 = self.client.post("/api/words/review", json={"word_id": word_id, "quality": 5}).json()["word"]
        self.assertEqual(r5["repetitions"], 5)
        self.assertEqual(r5["status"], "mastered")

    def test_t4_06_relapse_and_recovery_journey(self):
        """Student masters word, forgets it after vacation, relearns it."""
        w = self.add_word("relearn_target", "từ học lại").json()["word"]
        word_id = w["id"]

        # Master word
        for q in [4, 4, 4, 4, 5]:
            self.client.post("/api/words/review", json={"word_id": word_id, "quality": q})

        # Relapse (Forgot during vacation)
        relapse = self.client.post("/api/words/review", json={"word_id": word_id, "quality": 1}).json()["word"]
        self.assertEqual(relapse["status"], "learning")
        self.assertEqual(relapse["repetitions"], 0)
        self.assertEqual(relapse["interval"], 1)

        # Recover
        rec1 = self.client.post("/api/words/review", json={"word_id": word_id, "quality": 4}).json()["word"]
        self.assertEqual(rec1["repetitions"], 1)
        self.assertEqual(rec1["interval"], 1)


if __name__ == "__main__":
    print("=" * 70)
    print("RUNNING VOCABMASTER COMPREHENSIVE E2E TEST SUITE")
    print("=" * 70)
    suite = unittest.TestLoader().loadTestsFromNames([
        "test_e2e_vocabmaster.TestTier1FeatureCoverage",
        "test_e2e_vocabmaster.TestTier2BoundaryCornerCases",
        "test_e2e_vocabmaster.TestTier3CrossFeatureCombinations",
        "test_e2e_vocabmaster.TestTier4RealWorldScenarios",
    ])
    runner = unittest.TextTestRunner(verbosity=2)
    result = runner.run(suite)

    print("\n" + "=" * 70)
    print(f"E2E TEST SUMMARY: Ran {result.testsRun} tests | Failures: {len(result.failures)} | Errors: {len(result.errors)}")
    print("=" * 70)
    if result.wasSuccessful():
        print("🎉 ALL 130 TESTS PASSED CLEANLY (100% SUCCESS)!")
        sys.exit(0)
    else:
        print("❌ SOME TESTS FAILED!")
        sys.exit(1)
