"""
Automated Verification Suite for VocabMaster.
Tests file parsing (PDF, DOCX), vocabulary extraction, database CRUD, and SM-2 algorithm.
"""
import os
import sys

# Đảm bảo import được module app
sys.path.insert(0, os.path.dirname(__file__))

from app.parsers import extract_text_from_pdf, extract_text_from_docx, split_into_sentences
from app.vocab_engine import process_document_vocabulary
from app.srs import calculate_sm2
from app.database import init_db, save_document, save_words, get_all_words, update_srs_review, get_stats

def test_srs_algorithm():
    print("Testing SM-2 Algorithm...")
    # Lần 1 học tốt (quality 4) -> interval 1
    res1 = calculate_sm2(quality=4, repetitions=0, ease_factor=2.5, interval=1)
    assert res1["interval"] == 1, f"Expected 1, got {res1['interval']}"
    assert res1["repetitions"] == 1

    # Lần 2 học tốt (quality 4) -> interval 6
    res2 = calculate_sm2(quality=4, repetitions=1, ease_factor=res1["ease_factor"], interval=1)
    assert res2["interval"] == 6, f"Expected 6, got {res2['interval']}"
    assert res2["repetitions"] == 2

    # Lần 3 quên (quality 1) -> reset interval 1, repetitions 0
    res3 = calculate_sm2(quality=1, repetitions=2, ease_factor=res2["ease_factor"], interval=6)
    assert res3["interval"] == 1
    assert res3["repetitions"] == 0
    print("✓ SM-2 Algorithm passed successfully!")

def test_docx_parsing():
    print("\nTesting DOCX parsing with portable sample support...")
    # First verify dynamically generated valid DOCX file for 100% environment portability
    import io
    import docx
    doc = docx.Document()
    doc.add_heading("VocabMaster Study Guide: English Vocabulary", level=1)
    doc.add_paragraph("Resilience, perseverance, and dedication are vital qualities for students.")
    doc.add_paragraph("Healthy nutrition provides essential vitamins and minerals to maintain optimal wellness.")
    bio = io.BytesIO()
    doc.save(bio)
    dynamic_docx_bytes = bio.getvalue()
    
    dynamic_text = extract_text_from_docx(dynamic_docx_bytes)
    print(f"✓ Parsed portable dynamic DOCX: extracted {len(dynamic_text)} characters.")
    assert len(dynamic_text) > 0, "Extracted dynamic DOCX text is empty"
    assert "resilience" in dynamic_text.lower() or "qualities" in dynamic_text.lower()

    # If host sample exists, test it as well
    sample_docx = "/home/Huy/Downloads/BaoCao_DuAn_HoTro_NguoiKhiemThi.docx"
    if not os.path.exists(sample_docx):
        sample_docx = "/home/Huy/Downloads/Document1.docx"
    
    if os.path.exists(sample_docx):
        with open(sample_docx, "rb") as f:
            content = f.read()
        text = extract_text_from_docx(content)
        print(f"✓ Parsed host DOCX: {os.path.basename(sample_docx)}, extracted {len(text)} characters.")
        assert len(text) > 0, "Extracted text is empty"

def test_pdf_parsing_and_vocab_extraction():
    print("\nTesting PDF parsing and Vocabulary Extraction with portable sample support...")
    # First verify dynamically generated valid PDF file for 100% environment portability
    import pymupdf
    pdf_doc = pymupdf.open()
    page = pdf_doc.new_page()
    page.insert_text(
        (50, 72),
        "VocabMaster Study Material:\n"
        "Resilience, perseverance, and dedication are vital qualities for students.\n"
        "Healthy nutrition provides essential vitamins and minerals to maintain optimal wellness."
    )
    dynamic_pdf_bytes = pdf_doc.tobytes()
    pdf_doc.close()

    dynamic_text = extract_text_from_pdf(dynamic_pdf_bytes)
    print(f"✓ Parsed portable dynamic PDF: extracted {len(dynamic_text)} characters.")
    assert len(dynamic_text) > 0, "Extracted dynamic PDF text is empty"

    dynamic_words = process_document_vocabulary(dynamic_text, max_words=5)
    print(f"✓ Extracted {len(dynamic_words)} vocabulary words from portable dynamic PDF:")
    for w in dynamic_words[:3]:
        print(f"   - {w['word']} {w.get('ipa', '')}: {w['meaning_vi']}")
    assert len(dynamic_words) > 0, "Dynamic PDF vocabulary extraction failed"

    # If host sample exists, test it as well
    sample_pdf = "/home/Huy/Downloads/6. [TV-ĐHNT]_Kho từ vựng tuần 1 - Health.pdf"
    if not os.path.exists(sample_pdf):
        sample_pdf = "/home/Huy/Downloads/3000 từ tiếng Anh thông dụng nhất.pdf"

    if os.path.exists(sample_pdf):
        with open(sample_pdf, "rb") as f:
            content = f.read()
        text = extract_text_from_pdf(content)
        print(f"✓ Parsed host PDF: {os.path.basename(sample_pdf)}, extracted {len(text)} characters.")
        assert len(text) > 0, "Extracted text is empty"

def test_database_crud():
    print("\nTesting Database operations with isolated test DB...")
    import tempfile
    import app.database

    with tempfile.NamedTemporaryFile(suffix=".db", delete=False) as tmp:
        test_db_path = tmp.name

    original_db = app.database.DB_PATH
    try:
        app.database.DB_PATH = test_db_path
        init_db()
        stats_before = get_stats()
        assert stats_before["total_words"] == 0, "Test DB should be empty initially"

        saved = save_words(None, [{
            "word": "resilience",
            "pos": "n",
            "meaning_vi": "khả năng phục hồi, sự kiên cường",
            "ipa": "/rɪˈzɪl.jəns/",
            "example_sentence": "He showed great resilience during difficult times.",
            "example_vi": "Anh ấy đã thể hiện sự kiên cường tuyệt vời trong những thời điểm khó khăn."
        }])
        assert len(saved) == 1
        stats_after = get_stats()
        assert stats_after["total_words"] == 1, "DB should have 1 word"

        # Test review update
        words = get_all_words()
        first_word = words[0]
        updated = update_srs_review(first_word["id"], quality=4, mode="test")
        assert updated["repetitions"] == 1
        assert updated["status"] in ["learning", "mastered"]
        print(f"✓ Database CRUD & review verified for word: {updated['word']}")
    finally:
        app.database.DB_PATH = original_db
        try:
            os.remove(test_db_path)
        except Exception:
            pass

def test_thesaurus_engine():
    print("\nTesting Smart Thesaurus & Semantic Integrity Engine...")
    from app.vocab_engine import fetch_thesaurus
    result = fetch_thesaurus("abundant")
    assert result["word"] == "abundant"
    assert len(result["meaning_vi"]) > 0, "Meaning VI should not be empty"
    assert len(result["synonyms"]) >= 5, f"Expected >= 5 synonyms, got {len(result['synonyms'])}"
    assert len(result["antonyms"]) >= 1, f"Expected >= 1 antonyms, got {len(result['antonyms'])}"
    print(f"✓ Thesaurus verified for '{result['word']}': {result['meaning_vi']}")

    # Kiểm tra kiểm soát trôi dạt ngữ nghĩa (WordNet Semantic Drift Control)
    print("Testing semantic drift prevention on polysemous words...")
    
    # Test 1: 'sound' không được chứa 'liquid' hoặc 'gas' làm từ trái nghĩa
    res_sound = fetch_thesaurus("sound")
    ant_words_sound = {a["word"].lower() for a in res_sound.get("antonyms", [])}
    assert "liquid" not in ant_words_sound, f"WordNet drift: 'liquid' found in antonyms of 'sound': {ant_words_sound}"
    assert "gas" not in ant_words_sound, f"WordNet drift: 'gas' found in antonyms of 'sound': {ant_words_sound}"
    assert "unsound" in ant_words_sound or "unreliable" in ant_words_sound
    print("✓ 'sound': Passed (Zero liquid/gas in antonyms)")

    # Test 2: 'fair' không được chứa 'carnival' hoặc 'bazaar' làm từ đồng nghĩa
    res_fair = fetch_thesaurus("fair")
    syn_words_fair = {s["word"].lower() for s in res_fair.get("synonyms", [])}
    assert "carnival" not in syn_words_fair, f"POS drift: 'carnival' found in synonyms of 'fair': {syn_words_fair}"
    assert "bazaar" not in syn_words_fair, f"POS drift: 'bazaar' found in synonyms of 'fair': {syn_words_fair}"
    assert "just" in syn_words_fair or "equitable" in syn_words_fair
    print("✓ 'fair': Passed (Zero carnival/bazaar in synonyms)")

    # Test 3: 'fine' không được chứa 'mulct' hoặc 'amercement'
    res_fine = fetch_thesaurus("fine")
    syn_words_fine = {s["word"].lower() for s in res_fine.get("synonyms", [])}
    assert "mulct" not in syn_words_fine and "amercement" not in syn_words_fine
    print("✓ 'fine': Passed (Zero mulct/amercement)")

    # Test 4: 'lead' không được chứa 'tail' làm từ trái nghĩa
    res_lead = fetch_thesaurus("lead")
    ant_words_lead = {a["word"].lower() for a in res_lead.get("antonyms", [])}
    assert "tail" not in ant_words_lead, f"WordNet drift: 'tail' found in antonyms of 'lead': {ant_words_lead}"
    assert "follow" in ant_words_lead or "trail" in ant_words_lead
    print("✓ 'lead': Passed (Zero tail in antonyms)")

    # Test 5: 'patient' không được chứa 'outpatient' hoặc 'inpatient'
    res_pat = fetch_thesaurus("patient")
    ant_words_pat = {a["word"].lower() for a in res_pat.get("antonyms", [])}
    assert "outpatient" not in ant_words_pat and "inpatient" not in ant_words_pat
    assert "impatient" in ant_words_pat
    print("✓ 'patient': Passed (Zero outpatient in antonyms)")

    # Test 6: Ca biên cụm từ & hyphenated
    res_phr = fetch_thesaurus("well-known")
    assert len(res_phr["synonyms"]) >= 1
    print("✓ 'well-known': Edge case passed")

    # Test 7: 'noun' phải có nghĩa tiếng Việt chuẩn xác là 'danh từ' và POS là 'noun'
    res_noun = fetch_thesaurus("noun")
    assert res_noun["meaning_vi"].lower() != "noun", "Error: 'noun' meaning should be in Vietnamese, not English"
    assert "danh từ" in res_noun["meaning_vi"].lower()
    assert res_noun["pos"] == "noun", f"Expected pos 'noun', got '{res_noun['pos']}'"
    assert len(res_noun["synonyms"]) >= 5
    print("✓ 'noun': Meaning and POS verified (danh từ, pos=noun)")

if __name__ == "__main__":
    print("=====================================================")
    print("RUNNING VOCABMASTER TEST SUITE")
    print("=====================================================")
    test_srs_algorithm()
    test_docx_parsing()
    test_pdf_parsing_and_vocab_extraction()
    test_database_crud()
    test_thesaurus_engine()
    print("\n🎉 ALL TESTS PASSED SUCCESSFULLY!")

